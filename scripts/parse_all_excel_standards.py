import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

f1 = r"d:\FE-Projects\headcount-planning\.docs\Khung định biên nhân sự.xlsx"
f2 = r"d:\FE-Projects\headcount-planning\.docs\Xây dựng vòng đời chức danh PCD.xlsx"

wb1 = openpyxl.load_workbook(f1, data_only=True)
wb2 = openpyxl.load_workbook(f2, data_only=True)

print("================================================================================")
print("CHI TIẾT CẤU HÌNH TỪNG KHỐI TRONG 2 FILE EXCEL")
print("================================================================================")

# 1. PLP
print("\n--- 1. KHỐI PLP (Pháp lý dự án) ---")
s = wb1['PLP']
headers = [s.cell(5, c).value for c in range(5, 10)]
print("Tiêu chí quy mô:", headers)
for r in range(6, 12):
    role = s.cell(r, 4).value
    vals = [s.cell(r, c).value for c in range(5, 10)]
    print(f"Role: {role} | Values theo quy mô: {vals}")

# 2. DMD - Thấp tầng
print("\n--- 2. KHỐI DMD - NHÀ THẤP TẦNG ---")
s = wb1['DMD-NTT']
headers = [s.cell(6, c).value for c in range(5, 8)]
print("Tiêu chí quy mô:", headers)
for r in range(8, 18):
    dept = s.cell(r, 2).value
    role = s.cell(r, 4).value
    vals = [s.cell(r, c).value for c in range(5, 8)]
    note = s.cell(r, 8).value
    print(f"Role: {role} ({dept}) | Values: {vals} | Note: {note}")

# 3. DMD - Cao tầng
print("\n--- 3. KHỐI DMD - NHÀ CAO TẦNG ---")
s = wb1['DMD-NCT']
headers = [s.cell(6, c).value for c in range(5, 9)]
print("Tiêu chí quy mô:", headers)
for r in range(8, 18):
    dept = s.cell(r, 2).value
    role = s.cell(r, 4).value
    vals = [s.cell(r, c).value for c in range(5, 9)]
    note = s.cell(r, 9).value
    print(f"Role: {role} ({dept}) | Values: {vals} | Note: {note}")

# 4. PMD - Thấp tầng
print("\n--- 4. KHỐI PMD - THẤP TẦNG ---")
s = wb1['PMD.THẤP TẦNG']
headers = [s.cell(5, c).value for c in range(4, 8)]
print("Tiêu chí quy mô:", headers)
for r in range(6, 10):
    role = s.cell(r, 3).value
    vals = [s.cell(r, c).value for c in range(4, 8)]
    note = s.cell(r, 8).value
    print(f"Role: {role} | Values: {vals}")

# 5. PMD - Cao tầng
print("\n--- 5. KHỐI PMD - CAO TẦNG ---")
s = wb1['PMD.CAO TẦNG']
headers = [s.cell(5, c).value for c in range(4, 8)]
print("Tiêu chí quy mô:", headers)
for r in range(6, 10):
    role = s.cell(r, 3).value
    vals = [s.cell(r, c).value for c in range(4, 8)]
    print(f"Role: {role} | Values: {vals}")

# 6. PCD
print("\n--- 6. KHỐI PCD (Quản lý xây dựng) ---")
s = wb1['PCD.CỌC-HTKT-NHÀ THẤP TẦNG']
# Find stage blocks
for r in range(5, 75):
    giai_doan = s.cell(r, 2).value
    co_so = s.cell(r, 3).value
    role = s.cell(r, 5).value or s.cell(r, 4).value
    val1 = s.cell(r, 6).value
    val2 = s.cell(r, 7).value
    val3 = s.cell(r, 8).value
    if role and any([val1, val2, val3]):
        print(f"GĐ: {giai_doan or ''} | CS: {co_so or ''} | Role: {role} | Mức 1: {val1} | Mức 2: {val2} | Mức 3: {val3}")
