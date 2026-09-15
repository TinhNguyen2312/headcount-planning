import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

f1 = r"d:\FE-Projects\headcount-planning\.docs\Khung định biên nhân sự.xlsx"
f2 = r"d:\FE-Projects\headcount-planning\.docs\Xây dựng vòng đời chức danh PCD.xlsx"

print("================================================================================")
print("PHÂN TÍCH FILE 1: Khung định biên nhân sự.xlsx")
print("================================================================================")

wb1 = openpyxl.load_workbook(f1, data_only=True)

for sname in wb1.sheetnames:
    sheet = wb1[sname]
    print(f"\n>>> SHEET: [{sname}] (Rows: {sheet.max_row}, Cols: {sheet.max_column})")
    
    # Print non-empty header / data rows
    for r in range(1, min(35, sheet.max_row + 1)):
        row_vals = []
        for c in range(1, min(15, sheet.max_column + 1)):
            v = sheet.cell(r, c).value
            if v is not None:
                str_v = str(v).strip().replace('\n', ' ')
                if str_v:
                    row_vals.append(f"C{c}: {str_v}")
        if row_vals:
            print(f"  R{r:02d}: " + " | ".join(row_vals[:8]))

print("\n================================================================================")
print("PHÂN TÍCH FILE 2: Xây dựng vòng đời chức danh PCD.xlsx")
print("================================================================================")

wb2 = openpyxl.load_workbook(f2, data_only=True)

for sname in ['ND CHÍNH', 'KHUNG ĐB-THẤP TẦNG-PCD FULL', 'Hệ số tối ưu nhân lực', 'Tính toán hệ số', 'Modul 100 căn thô']:
    if sname not in wb2.sheetnames:
        continue
    sheet = wb2[sname]
    print(f"\n>>> SHEET: [{sname}] (Rows: {sheet.max_row}, Cols: {sheet.max_column})")
    for r in range(1, min(35, sheet.max_row + 1)):
        row_vals = []
        for c in range(1, min(15, sheet.max_column + 1)):
            v = sheet.cell(r, c).value
            if v is not None:
                str_v = str(v).strip().replace('\n', ' ')
                if str_v:
                    row_vals.append(f"C{c}: {str_v}")
        if row_vals:
            print(f"  R{r:02d}: " + " | ".join(row_vals[:8]))
